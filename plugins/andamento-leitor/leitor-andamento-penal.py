from docx import Document
from openpyxl import Workbook 
import re

document = Document('ANDAMENTO_PENAL_MODIFICADO.docx')

iLinha = 1

arquivo_excel = Workbook()
planilha2 = arquivo_excel.create_sheet("PECs")
planilha2['A1'] = 'IPEN'
planilha2['B1'] = 'PEC'

for temp in document.paragraphs:
    positionIpen = temp.text.find('ipen')
    
    if positionIpen  != -1:
        iLinha += 1

        planilha2['A' + str(iLinha)] = temp.text[(positionIpen + 4) : (positionIpen + 11)]
        planilha2['B' + str(iLinha)] = temp.text[(positionIpen + 18) : (positionIpen + 43)]

        print(temp.text[(positionIpen + 18) : (positionIpen + 43)])

arquivo_excel.save("pec.xlsx")